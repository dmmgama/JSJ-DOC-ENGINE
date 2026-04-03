"""
preprocessor.py — JSJ-DOC-ENGINE
Substitui tags {{ excel | path | sheet | range }} em Markdown
por tabelas Markdown geradas a partir de ficheiros Excel via openpyxl.

Interface pública:
    process_markdown(md_text: str, base_path: str | Path) -> str
"""

import re
from pathlib import Path

import openpyxl

# Regex para tags: {{ excel | path/to/file.xlsx | SheetName | A1:D10 }}
_TAG_PATTERN = re.compile(
    r'\{\{\s*excel\s*\|\s*([^|]+?)\s*\|\s*([^|]+?)\s*\|\s*([^}]+?)\s*\}\}'
)


def _parse_range(range_str: str):
    """Convert 'A1:D10' into (min_row, min_col, max_row, max_col) 1-based ints."""
    parts = range_str.upper().split(':')
    if len(parts) != 2:
        raise ValueError(f"Intervalo inválido: '{range_str}'. Formato esperado: A1:D10")
    start, end = parts

    def cell_to_rc(cell_ref):
        match = re.fullmatch(r'([A-Z]+)(\d+)', cell_ref)
        if not match:
            raise ValueError(f"Referência de célula inválida: '{cell_ref}'")
        col_str, row_str = match.group(1), match.group(2)
        col = 0
        for ch in col_str:
            col = col * 26 + (ord(ch) - ord('A') + 1)
        return int(row_str), col

    r1, c1 = cell_to_rc(start)
    r2, c2 = cell_to_rc(end)
    return r1, c1, r2, c2


def _read_excel_range(excel_path: Path, sheet_name: str, range_str: str) -> list[list]:
    """Lê um range de um ficheiro Excel e devolve lista de listas de valores."""
    if not excel_path.exists():
        raise FileNotFoundError(
            f"Ficheiro Excel não encontrado: '{excel_path}'\n"
            "Verifique o path na tag {{ excel | ... }}"
        )

    try:
        wb = openpyxl.load_workbook(str(excel_path), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f"Erro ao abrir '{excel_path}': {exc}") from exc

    if sheet_name not in wb.sheetnames:
        available = ', '.join(wb.sheetnames)
        raise ValueError(
            f"Sheet '{sheet_name}' não encontrada em '{excel_path.name}'. "
            f"Sheets disponíveis: {available}"
        )

    ws = wb[sheet_name]
    min_row, min_col, max_row, max_col = _parse_range(range_str)

    rows = []
    for row in ws.iter_rows(
        min_row=min_row, max_row=max_row,
        min_col=min_col, max_col=max_col,
        values_only=True,
    ):
        rows.append([str(cell) if cell is not None else '' for cell in row])

    wb.close()
    return rows


def _rows_to_markdown_table(rows: list[list]) -> str:
    """Converte lista de listas numa tabela Markdown (primeira linha = cabeçalho)."""
    if not rows:
        return ''

    header = rows[0]
    body = rows[1:]

    def _escape(cell: str) -> str:
        return cell.replace('|', '\\|')

    def _row_str(cells):
        return '| ' + ' | '.join(_escape(c) for c in cells) + ' |'

    separator = '| ' + ' | '.join(['---'] * len(header)) + ' |'

    lines = [_row_str(header), separator]
    for row in body:
        # Garante que a linha tem o mesmo número de colunas que o cabeçalho
        padded = row + [''] * (len(header) - len(row))
        lines.append(_row_str(padded[:len(header)]))

    return '\n'.join(lines)


def process_markdown(md_text: str, base_path) -> str:
    """
    Substitui todas as tags {{ excel | path | sheet | range }} no texto MD
    por tabelas Markdown geradas a partir dos ficheiros Excel correspondentes.

    Args:
        md_text:    Texto Markdown a processar.
        base_path:  Directório base para resolver paths relativos das tags.

    Returns:
        Texto Markdown com as tags substituídas por tabelas.
        Se não houver tags, devolve md_text inalterado.

    Raises:
        FileNotFoundError: Se um ficheiro Excel referenciado não existir.
        ValueError: Se a tag for malformada ou a sheet/range for inválida.
    """
    base_path = Path(base_path)

    def _replace(match):
        raw_path = match.group(1).strip()
        sheet_name = match.group(2).strip()
        range_str = match.group(3).strip()

        excel_path = (base_path / raw_path).resolve()
        rows = _read_excel_range(excel_path, sheet_name, range_str)

        if not rows:
            return f'*(tabela vazia: {raw_path} | {sheet_name} | {range_str})*'

        return _rows_to_markdown_table(rows)

    return _TAG_PATTERN.sub(_replace, md_text)


# ---------------------------------------------------------------------------
# Smoke test — executa apenas quando invocado directamente
# ---------------------------------------------------------------------------
if __name__ == '__main__':
    sample_md = "# Título\n\nParágrafo sem tags Excel.\n\n## Secção 2\n\nTexto normal.\n"
    result = process_markdown(sample_md, base_path='.')
    assert result == sample_md, "FALHOU: MD sem tags foi alterado"
    print("OK — smoke test passou: MD sem tags devolvido inalterado.")
